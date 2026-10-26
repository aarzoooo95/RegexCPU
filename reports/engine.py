"""Word-report engine: reproduces the layout of the Alliance sample report
(bordered pages, cover, certificate, declaration, acknowledgement, TOC, chapters)."""
import os, re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
TEXT_WIDTH_IN = 5.9


def _font(run, name=FONT, size=None, bold=None, italic=None, underline=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if underline is not None: run.font.underline = underline
    if color: run.font.color.rgb = RGBColor.from_string(color)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig = self.tab = 0
        self.toc_marks = []                 # (text, level)
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        s.left_margin = s.right_margin = Inches(1.18)
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(0.9)
        self._border(s)
        self._footer(s)
        st = d.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        for name, size, align in (("Heading 1", 18, AL.CENTER), ("Heading 2", 14, AL.LEFT), ("Heading 3", 12, AL.LEFT)):
            h = d.styles[name]
            h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(size), True, False
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.element.rPr.rFonts.set(qn("w:ascii"), FONT); h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            h.paragraph_format.alignment = align
            h.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 8)
            h.paragraph_format.space_after = Pt(8 if name == "Heading 1" else 6)
            h.paragraph_format.keep_with_next = True

    # ---- page furniture -------------------------------------------------
    def _border(self, section):
        pg = OxmlElement("w:pgBorders"); pg.set(qn("w:offsetFrom"), "page")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            for k, v in (("val", "single"), ("sz", "12"), ("space", "24"), ("color", "000000")):
                el.set(qn(f"w:{k}"), v)
            pg.append(el)
        section._sectPr.find(qn("w:pgMar")).addnext(pg)

    def _footer(self, section):
        p = section.footer.paragraphs[0]; p.alignment = AL.CENTER
        r = p.add_run(); _font(r, size=10)
        for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
            r._element.append(el)

    # ---- text -----------------------------------------------------------
    def p(self, text="", size=12, bold=False, italic=False, align=None, after=6, before=0,
          line=1.5, underline=False, keep_next=False, indent=None, style=None):
        para = self.doc.add_paragraph(style=style)
        pf = para.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), line
        if align is not None: para.alignment = align
        if keep_next: pf.keep_with_next = True
        if indent is not None: pf.left_indent = Inches(indent)
        self.rich(para, text, size, bold, italic, underline)
        return para

    def rich(self, para, text, size=12, bold=False, italic=False, underline=False):
        """**bold**  and  `code`  inline markup."""
        for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
            if not part: continue
            if part.startswith("**"):
                _font(para.add_run(part[2:-2]), size=size, bold=True, italic=italic, underline=underline)
            elif part.startswith("`"):
                _font(para.add_run(part[1:-1]), name=MONO, size=size - 1.5, bold=bold)
            else:
                _font(para.add_run(part), size=size, bold=bold, italic=italic, underline=underline)

    def body(self, text):
        return self.p(text, align=AL.JUSTIFY)

    def bullets(self, items, size=12):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(3); para.paragraph_format.line_spacing = 1.35
            para.alignment = AL.JUSTIFY
            self.rich(para, it, size)

    def numbered(self, items, size=12):
        for it in items:
            para = self.doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(3); para.paragraph_format.line_spacing = 1.35
            self.rich(para, it, size)

    def h1(self, text, new_page=True):
        h = self.doc.add_heading(text, level=1)
        h.paragraph_format.page_break_before = new_page
        self.toc_marks.append((text, 0)); return h

    def h2(self, text, toc=True):
        h = self.doc.add_heading(text, level=2)
        if toc: self.toc_marks.append((text, 1))
        return h

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- code / tables / figures -----------------------------------------
    def code(self, text, size=8.5):
        for line in text.rstrip("\n").split("\n"):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = pf.space_before = Pt(0); pf.line_spacing = 1.0
            pf.left_indent = Inches(0.08)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
            para._p.get_or_add_pPr().append(shd)
            _font(para.add_run(line.replace("\t", "    ") or " "), name=MONO, size=size)
        self.p("", after=4, line=1.0)

    def table(self, header, rows, widths=None, size=10, caption=None, align_first_left=True):
        if caption:
            self.tab += 1
            self.p(f"Table {self.tab}: {caption}", size=10, bold=True, align=AL.CENTER, after=3, line=1.0, keep_next=True)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""
            self._cell(c, str(h), size, True, AL.CENTER, "D9E2F3")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                self._cell(cells[i], str(v), size, False, AL.LEFT if (i == 0 and align_first_left) else AL.CENTER)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.p("", after=4, line=1.0)
        return t

    def _cell(self, cell, text, size, bold, align, fill=None):
        para = cell.paragraphs[0]; para.alignment = align
        para.paragraph_format.space_after = Pt(1); para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.line_spacing = 1.0
        self.rich(para, text, size, bold)
        if fill:
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
            cell._tc.get_or_add_tcPr().append(shd)

    def image(self, path, width, caption=None):
        para = self.doc.add_paragraph(); para.alignment = AL.CENTER
        para.paragraph_format.keep_with_next = bool(caption); para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(path, width=Inches(width))
        if caption:
            self.fig += 1
            self.p(f"Figure {self.fig}: {caption}", size=10, italic=True, align=AL.CENTER, after=8, line=1.0)

    # ---- fixed front matter ----------------------------------------------
    def cover(self, cfg, title, subject, guide, assets):
        self.p("", after=18)
        self.p("MINI PROJECT REPORT", size=16, bold=True, align=AL.CENTER, after=20, line=1.0)
        self.p(title.upper(), size=12, bold=True, italic=True, align=AL.CENTER, after=18, line=1.15)
        self.p("Submitted by:", bold=True, align=AL.CENTER, after=0, line=1.15)
        self.p(cfg["name"], align=AL.CENTER, after=0, line=1.15)
        self.p(cfg["roll_no"], align=AL.CENTER, after=14, line=1.15)
        self.p("In partial fulfilment for the award of the degree of", italic=True, align=AL.CENTER, after=10, line=1.15)
        self.p(cfg["degree"], bold=True, align=AL.CENTER, after=0, line=1.15)
        self.p("In", align=AL.CENTER, after=0, line=1.15)
        self.p(cfg["department"], bold=True, align=AL.CENTER, after=14, line=1.15)
        self.p("Under the guidance of", align=AL.CENTER, after=0, line=1.15)
        self.p(guide, bold=True, align=AL.CENTER, after=4, line=1.15)
        self.p(cfg["designation"], bold=True, align=AL.CENTER, after=18, line=1.15)
        para = self.doc.add_paragraph(); para.alignment = AL.CENTER; para.paragraph_format.space_after = Pt(8)
        para.add_run().add_picture(os.path.join(assets, "logo_cover.png"), width=Inches(2.6), height=Inches(2.6))
        self.p("Alliance School Of Advance Computing", size=16, bold=True, align=AL.CENTER, after=8, before=6, line=1.15)
        self.p("Alliance University, Bengaluru", size=14, bold=True, align=AL.CENTER, after=8, line=1.15)
        self.p(cfg["month_year"], size=14, bold=True, align=AL.CENTER, after=0, line=1.15)

    def _school_header(self, assets, logo_w=1.0):
        self.page_break()
        self.p("", after=40)
        self.image(os.path.join(assets, "logo_small"+".png"), logo_w)
        self.p("", after=14)
        self.p("COMPUTER SCIENCE AND ENGINEERING", size=14, align=AL.CENTER, after=0, line=1.15)
        self.p("ALLIANCE SCHOOL OF ADVANCED COMPUTING", size=14, align=AL.CENTER, after=16, line=1.15)

    def certificate(self, cfg, topic, guide, assets):
        self._school_header(assets)
        self.p("CERTIFICATE", size=14, underline=True, align=AL.CENTER, after=10, line=1.15)
        self.p(f"This is to certify that the report on {topic} is submitted by {cfg['name']} "
               f"[Uni. Roll No. {cfg['roll_no']}] in partial fulfilment for the award of the degree of "
               f"{cfg['degree'].title()} (Computer Science and Engineering) of Alliance University, is a bonafide "
               f"work accomplished under my supervision and guidance during the academic year {cfg['academic_year']}. "
               "This report embodies the results of original work and studies conducted by students and the "
               "contents do not form the basis for the award of any other degree to the candidate or anybody else.",
               size=12, align=AL.JUSTIFY, line=1.5, after=36)
        self.p("----------------------", align=AL.RIGHT, after=0, line=1.0)
        self.p(guide, align=AL.RIGHT, after=0, line=1.15)
        self.p(cfg["designation"].split(",")[0], align=AL.RIGHT, after=0, line=1.15)

    def declaration(self, cfg, title, guide, assets):
        self._school_header(assets)
        self.p("DECLARATION", size=14, underline=True, align=AL.CENTER, after=14, line=1.15)
        para = self.p("", align=AL.JUSTIFY, after=6)
        self.rich(para, f"I hereby declare that the report titled {title.upper()}, submitted by me in partial fulfilment of "
                  f"the requirements for the award of the degree of **{cfg['degree'].title()} in Computer Science and "
                  f"Engineering (CSE)** at **Alliance University**, is a record of my original work carried out under the "
                  f"supervision of **{guide}, {cfg['designation']}**.")
        self.p("I further confirm that all views, findings, and conclusions expressed herein are the result of sincere "
               "effort and academic inquiry, and have been reviewed and deliberated in consultation with the assigned mentor.",
               align=AL.JUSTIFY, after=6)
        self.p("This declaration stands as a testament to the authenticity and integrity of the work submitted.",
               align=AL.JUSTIFY, after=48)
        for a, b in ((cfg["name"], cfg["roll_no"]), ("Name of the Student", "University Registration Number")):
            para = self.doc.add_paragraph(); para.paragraph_format.space_after = Pt(10)
            ts = para.paragraph_format.tab_stops
            ts.add_tab_stop(Inches(2.5)); ts.add_tab_stop(Inches(5.0))
            small = a == "Name of the Student"
            _font(para.add_run(f"{a}\t{b}" + ("\tSignature" if small else "")), size=10 if small else 12)

    def acknowledgement(self, cfg, title, guide, assets, closing):
        self.page_break()
        self.p("Acknowledgement", size=16, bold=True, align=AL.LEFT, after=10, before=6, line=1.15)
        self.p(f"I would like to express my heartfelt gratitude to everyone who contributed to the successful completion "
               f"of this project titled {title.upper()}.", align=AL.JUSTIFY)
        self.p(f"First and foremost, I extend my sincere thanks to my mentor and guide, {guide}, for the support, insightful "
               "feedback, and encouragement throughout the development of this work. Their expertise and commitment to "
               "fostering learning provided the foundation for this project.", align=AL.JUSTIFY)
        self.p("I am deeply grateful to my faculty members and peers whose guidance and technical inputs helped refine the "
               "project. Their collaborative spirit and constructive suggestions were instrumental in shaping the final outcome.",
               align=AL.JUSTIFY)
        self.p(closing, align=AL.JUSTIFY)
        self.p("Lastly, I would like to thank everyone who supported me directly or indirectly during this journey.", align=AL.JUSTIFY)
        self.p("Thank you.", after=2)
        self.p(cfg["name"], bold=False)

    # ---- table of contents -------------------------------------------------
    def toc_placeholder(self):
        self.page_break()
        self.p("TABLE OF CONTENTS", size=14, bold=True, align=AL.LEFT, after=8, before=6, line=1.15)
        self._toc_anchor = self.doc.add_paragraph()

    def fill_toc(self, entries, pages):
        """entries: [(text, level)] in order; pages: {text: page}"""
        anchor = self._toc_anchor._p
        for text, level in entries:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(3); pf.line_spacing = 1.3; pf.left_indent = Inches(0.3 * level)
            pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            _font(para.add_run(f"{text}\t{pages.get(text, '')}"), size=12, bold=(level == 0))
            anchor.addprevious(para._p)
        anchor.getparent().remove(anchor)

    def save(self, path):
        self.doc.save(path)
